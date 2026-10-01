from playwright.sync_api import sync_playwright
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime
from pathlib import Path
import json
import random
import re
import time

BASE_DIR = Path(__file__).resolve().parent
WHATSAPP_URL = "https://web.whatsapp.com/"
CONTACTS_FILE = BASE_DIR / "contacts.xlsx"
PROFILE_DIR = BASE_DIR / "whatsapp_profile"
SCREENSHOT_DIR = BASE_DIR / "whatsapp_screenshots"
TODAY = datetime.now().strftime("%Y-%m-%d")
JSON_REPORT = BASE_DIR / f"whatsapp_report_{TODAY}.json"
EXCEL_REPORT = BASE_DIR / f"whatsapp_report_{TODAY}.xlsx"
DEFAULT_MESSAGE = "Hello {name}, hope you are doing well."

SCREENSHOT_DIR.mkdir(parents=True, exist_ok=True)


def human_pause(page, minimum=2, maximum=5):
    page.wait_for_timeout(random.randint(minimum * 1000, maximum * 1000))


def clean_filename(value):
    return re.sub(r'[<>:"/\\|?*]', "_", str(value)).strip()


def normalize_phone(phone):
    phone = str(phone or "").strip()
    if not phone:
        return ""
    if phone.startswith("+"):
        return "+" + re.sub(r"\D", "", phone[1:])
    return re.sub(r"\D", "", phone)


def read_contacts():
    if not CONTACTS_FILE.exists():
        raise FileNotFoundError(
            f"contacts.xlsx was not found. Expected location: {CONTACTS_FILE}"
        )

    workbook = load_workbook(CONTACTS_FILE, data_only=True)
    sheet = workbook.active

    headers = {}
    for column in range(1, sheet.max_column + 1):
        value = sheet.cell(row=1, column=column).value
        if value:
            headers[str(value).strip().lower()] = column

    for required in ["name", "phone"]:
        if required not in headers:
            raise ValueError(f"Missing required column: {required.title()}")

    contacts = []

    for row in range(2, sheet.max_row + 1):
        name = sheet.cell(row=row, column=headers["name"]).value
        phone = sheet.cell(row=row, column=headers["phone"]).value

        message = None
        if "message" in headers:
            message = sheet.cell(
                row=row,
                column=headers["message"]
            ).value

        if not name and not phone:
            continue

        contacts.append({
            "name": str(name or "").strip(),
            "phone": normalize_phone(phone),
            "message": str(message or DEFAULT_MESSAGE).strip()
        })

    workbook.close()
    return contacts


def find_search_box(page):
    print("Looking for WhatsApp search box...")

    selectors = [
        '#side div[contenteditable="true"][role="textbox"]',
        '#side [contenteditable="true"][aria-label*="Search"]',
        '#side [contenteditable="true"][data-tab="3"]',
        'div[contenteditable="true"][aria-label*="Search"]',
        '[role="textbox"][aria-label*="Search"]',
        '[contenteditable="true"][data-tab="3"]'
    ]