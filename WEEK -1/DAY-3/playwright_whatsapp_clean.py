from playwright.sync_api import sync_playwright
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime
from pathlib import Path
import json
import random
import re
import time

BASE_DIR = Path(__file__).resolve().parent
WHATSAPP_URL = "https://web.whatsapp.com/"
CONTACTS_FILE = BASE_DIR / "contacts.xlsx"
PROFILE_DIR = BASE_DIR / "whatsapp_profile"
SCREENSHOT_DIR = BASE_DIR / "whatsapp_screenshots"
TODAY = datetime.now().strftime("%Y-%m-%d")
JSON_REPORT = BASE_DIR / f"whatsapp_report_{TODAY}.json"
EXCEL_REPORT = BASE_DIR / f"whatsapp_report_{TODAY}.xlsx"
DEFAULT_MESSAGE = "Hello {name}, hope you are doing well."

SCREENSHOT_DIR.mkdir(parents=True, exist_ok=True)


def human_pause(page, minimum=2, maximum=5):
    page.wait_for_timeout(random.randint(minimum * 1000, maximum * 1000))


def clean_filename(value):
    return re.sub(r'[<>:"/\\|?*]', "_", str(value)).strip()


def normalize_phone(phone):
    phone = str(phone or "").strip()
    if not phone:
        return ""
    if phone.startswith("+"):
        return "+" + re.sub(r"\D", "", phone[1:])
    return re.sub(r"\D", "", phone)


def read_contacts():
    if not CONTACTS_FILE.exists():
        raise FileNotFoundError(
            f"contacts.xlsx was not found. Expected location: {CONTACTS_FILE}"
        )

    workbook = load_workbook(CONTACTS_FILE, data_only=True)
    sheet = workbook.active

    headers = {}
    for column in range(1, sheet.max_column + 1):
        value = sheet.cell(row=1, column=column).value
        if value:
            headers[str(value).strip().lower()] = column

    for required in ["name", "phone"]:
        if required not in headers:
            raise ValueError(f"Missing required column: {required.title()}")

    contacts = []

    for row in range(2, sheet.max_row + 1):
        name = sheet.cell(row=row, column=headers["name"]).value
        phone = sheet.cell(row=row, column=headers["phone"]).value

        message = None
        if "message" in headers:
            message = sheet.cell(
                row=row,
                column=headers["message"]
            ).value

        if not name and not phone:
            continue

        contacts.append({
            "name": str(name or "").strip(),
            "phone": normalize_phone(phone),
            "message": str(message or DEFAULT_MESSAGE).strip()
        })

    workbook.close()
    return contacts


def find_search_box(page):
    print("Looking for WhatsApp search box...")

    selectors = [
        '#side div[contenteditable="true"][role="textbox"]',
        '#side [contenteditable="true"][aria-label*="Search"]',
        '#side [contenteditable="true"][data-tab="3"]',
        'div[contenteditable="true"][aria-label*="Search"]',
        '[role="textbox"][aria-label*="Search"]',
        '[contenteditable="true"][data-tab="3"]'
    ]

    for selector in selectors:
        try:
            locator = page.locator(selector)
            for i in range(locator.count()):
                candidate = locator.nth(i)
                if candidate.is_visible():
                    print(f"Search box found using selector: {selector}")
                    return candidate
        except Exception:
            pass

    try:
        placeholder = page.get_by_text(
            "Search or start a new chat",
            exact=False
        )

        for i in range(placeholder.count()):
            item = placeholder.nth(i)

            if not item.is_visible():
                continue

            item.click()
            page.wait_for_timeout(500)

            focused = page.locator(":focus")

            if focused.count() > 0 and focused.first.is_visible():
                print("Search box found through focused element.")
                return focused.first
    except Exception:
        pass

    try:
        editables = page.locator('#side [contenteditable="true"]')
        candidates = []

        for i in range(editables.count()):
            item = editables.nth(i)

            if not item.is_visible():
                continue

            box = item.bounding_box()

            if not box:
                continue

            aria = item.get_attribute("aria-label") or ""
            role = item.get_attribute("role") or ""
            data_tab = item.get_attribute("data-tab") or ""

            print(
                f"Editable {i}: aria={aria!r}, role={role!r}, "
                f"data-tab={data_tab!r}, x={box['x']:.0f}, "
                f"y={box['y']:.0f}, w={box['width']:.0f}, h={box['height']:.0f}"
            )

            if box["width"] > 150 and box["height"] > 20 and box["y"] < 450:
                candidates.append((box["y"], item))

        if candidates:
            candidates.sort(key=lambda x: x[0])
            print("Search box selected using position-based fallback.")
            return candidates[0][1]

    except Exception:
        pass

    raise RuntimeError("WhatsApp search box could not be located.")


def find_message_box(page):
    selectors = [
        'footer div[contenteditable="true"][aria-label*="Type a message"]',
        'footer div[role="textbox"][contenteditable="true"]',
        'footer div[contenteditable="true"]',
        'div[contenteditable="true"][aria-label*="message"]'
    ]

    for selector in selectors:
        try:
            locator = page.locator(selector)

            for i in range(locator.count() - 1, -1, -1):
                candidate = locator.nth(i)

                if candidate.is_visible():
                    return candidate
        except Exception:
            pass

    raise RuntimeError("WhatsApp message box could not be located.")


def wait_for_whatsapp_login(page):
    print("\nOpening WhatsApp Web...")

    page.goto(
        WHATSAPP_URL,
        wait_until="domcontentloaded",
        timeout=60000
    )

    print("\nIf a QR code is shown, scan it with your phone.")
    print("Waiting for WhatsApp interface...")

    deadline = time.time() + 300

    while time.time() < deadline:
        try:
            search_box = find_search_box(page)

            if search_box.is_visible():
                print("\nWhatsApp login detected successfully!")
                page.wait_for_timeout(2000)
                return
        except Exception:
            pass

        try:
            side_panel = page.locator("#side")

            if side_panel.count() > 0 and side_panel.first.is_visible():
                print("\nWhatsApp main interface detected.")
                page.wait_for_timeout(2000)
                return
        except Exception:
            pass

        print("Still waiting for WhatsApp...")
        page.wait_for_timeout(2000)

    raise RuntimeError("WhatsApp login could not be detected within 5 minutes.")


def clear_search_box(search_box):
    search_box.click()
    time.sleep(0.5)
    search_box.press("Control+A")
    search_box.press("Backspace")
    time.sleep(0.5)


def try_open_contact(page, query):
    query = str(query or "").strip()

    if not query:
        return False

    print(f"Trying search: {query}")

    search_box = find_search_box(page)
    print("Search box located successfully.")

    clear_search_box(search_box)

    try:
        search_box.click()
        search_box.fill(query)
    except Exception:
        search_box.click()
        page.keyboard.type(query, delay=80)

    human_pause(page, 2, 4)

    try:
        no_results = page.get_by_text("No results found", exact=False)

        if no_results.count() > 0 and no_results.first.is_visible():
            return False
    except Exception:
        pass

    try:
        exact = page.get_by_text(query, exact=True)

        for i in range(exact.count()):
            candidate = exact.nth(i)

            if candidate.is_visible():
                candidate.click(timeout=5000)
                page.wait_for_timeout(2000)
                find_message_box(page)
                return True
    except Exception:
        pass

    try:
        titled = page.locator("span[title]")
        search_digits = re.sub(r"\D", "", query)

        for i in range(titled.count()):
            candidate = titled.nth(i)

            if not candidate.is_visible():
                continue

            title = candidate.get_attribute("title") or ""
            title_digits = re.sub(r"\D", "", title)

            if (
                query.lower() == title.lower()
                or (search_digits and search_digits == title_digits)
            ):
                candidate.click(timeout=5000)
                page.wait_for_timeout(2000)
                find_message_box(page)
                return True
    except Exception:
        pass

    for selector in [
        '#pane-side div[role="listitem"]',
        '#pane-side div[role="row"]',
        '#pane-side div[tabindex="-1"]'
    ]:
        try:
            rows = page.locator(selector)

            for i in range(rows.count()):
                row = rows.nth(i)

                if not row.is_visible():
                    continue

                try:
                    row.click(timeout=5000)
                    page.wait_for_timeout(2000)
                    find_message_box(page)
                    return True
                except Exception:
                    continue
        except Exception:
            pass

    return False


def open_contact(page, name, phone):
    print(f"\nSearching for contact:\nName : {name}\nPhone: {phone}")

    if name and try_open_contact(page, name):
        return True, "name"

    if phone and try_open_contact(page, phone):
        return True, "phone"

    return False, None


def build_personalized_message(name, template):
    return template.replace("{name}", name)


def send_message(page, message):
    message_box = find_message_box(page)
    outgoing_before = page.locator("div.message-out").count()

    message_box.click()
    human_pause(page, 2, 3)

    try:
        message_box.fill(message)
    except Exception:
        message_box.press_sequentially(
            message,
            delay=random.randint(30, 80)
        )

    human_pause(page, 2, 4)
    message_box.press("Enter")

    print("Message submitted.")

    deadline = time.time() + 20

    while time.time() < deadline:
        outgoing_after = page.locator("div.message-out").count()

        if outgoing_after > outgoing_before:
            print("Outgoing message detected.")
            page.wait_for_timeout(1500)
            return True

        page.wait_for_timeout(1000)

    return False


def extract_last_three_incoming(page):
    incoming = page.locator("div.message-in")
    count = incoming.count()

    if count == 0:
        return []

    extracted = []
    start = max(0, count - 10)

    for index in range(start, count):
        container = incoming.nth(index)

        try:
            text_locator = container.locator("span.selectable-text")

            if text_locator.count() > 0:
                pieces = []

                for j in range(text_locator.count()):
                    piece = text_locator.nth(j).inner_text(timeout=2000).strip()

                    if piece:
                        pieces.append(piece)

                text = " ".join(pieces)
            else:
                text = container.inner_text(timeout=2000).strip()

            if text:
                extracted.append(text)

        except Exception:
            continue

    return extracted[-3:]


def save_chat_screenshot(page, contact_name, index):
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_name = clean_filename(contact_name) or "contact"

    filename = f"{index:02d}_{safe_name}_{timestamp}.png"
    filepath = SCREENSHOT_DIR / filename

    page.screenshot(
        path=str(filepath),
        full_page=False
    )

    return str(filepath)


def save_json_report(results):
    report = {
        "report_date": TODAY,
        "generated_at": datetime.now().isoformat(),
        "total_contacts": len(results),
        "sent": sum(
            1 for item in results
            if item.get("status") == "SENT"
        ),
        "results": results
    }

    with open(JSON_REPORT, "w", encoding="utf-8") as file:
        json.dump(
            report,
            file,
            indent=4,
            ensure_ascii=False
        )


def save_excel_report(results):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "WhatsApp Report"

    headers = [
        "Name",
        "Phone",
        "Sent Message",
        "Status",
        "Search Method",
        "Last Incoming Message 1",
        "Last Incoming Message 2",
        "Last Incoming Message 3",
        "Screenshot",
        "Error",
        "Timestamp"
    ]

    sheet.append(headers)

    header_fill = PatternFill(
        "solid",
        fgColor="1F4E78"
    )

    header_font = Font(
        bold=True,
        color="FFFFFF"
    )

    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    for result in results:
        incoming = list(
            result.get(
                "last_3_incoming_messages",
                []
            )
        )

        while len(incoming) < 3:
            incoming.append("")

        sheet.append([
            result.get("name", ""),
            result.get("phone", ""),
            result.get("sent_message", ""),
            result.get("status", ""),
            result.get("search_method", ""),
            incoming[0],
            incoming[1],
            incoming[2],
            result.get("screenshot", ""),
            result.get("error", ""),
            result.get("timestamp", "")
        ])

    widths = {
        "A": 20,
        "B": 20,
        "C": 50,
        "D": 20,
        "E": 18,
        "F": 40,
        "G": 40,
        "H": 40,
        "I": 55,
        "J": 50,
        "K": 24
    }

    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )

    for row in range(2, sheet.max_row + 1):
        sheet[f"B{row}"].number_format = "@"

    sheet.freeze_panes = "A2"
    workbook.save(EXCEL_REPORT)


def process_contact(page, contact, index):
    name = contact["name"]
    phone = contact["phone"]
    template = contact["message"]

    personalized = build_personalized_message(
        name,
        template
    )

    result = {
        "name": name,
        "phone": phone,
        "sent_message": personalized,
        "status": "FAILED",
        "search_method": None,
        "last_3_incoming_messages": [],
        "screenshot": None,
        "error": None,
        "timestamp": datetime.now().isoformat()
    }

    try:
        print("\n" + "=" * 65)
        print(f"CONTACT {index}: {name or phone}")
        print("=" * 65)

        found, search_method = open_contact(
            page,
            name,
            phone
        )

        if not found:
            result["status"] = "CONTACT_NOT_FOUND"
            result["error"] = (
                "Contact could not be found using name or phone."
            )
            return result

        result["search_method"] = search_method

        print(f"Chat opened using: {search_method}")

        human_pause(page, 2, 4)

        print("\nPersonalized message:")
        print(personalized)

        confirmed = send_message(
            page,
            personalized
        )

        if confirmed:
            result["status"] = "SENT"
        else:
            result["status"] = "SENT_UNCONFIRMED"

        human_pause(page, 2, 4)

        incoming_messages = extract_last_three_incoming(
            page
        )

        result["last_3_incoming_messages"] = incoming_messages

        screenshot = save_chat_screenshot(
            page,
            name or phone,
            index
        )

        result["screenshot"] = screenshot

    except Exception as error:
        result["status"] = "FAILED"
        result["error"] = str(error)

        print(
            f"ERROR processing {name or phone}: {error}"
        )

    return result


def main():
    print("=" * 70)
    print("WHATSAPP MESSAGE SENDER + SMART DATA EXTRACTOR")
    print("=" * 70)

    print(f"\nScript folder:\n{BASE_DIR}")
    print(f"\nContacts file:\n{CONTACTS_FILE}")

    contacts = read_contacts()

    print(f"\nContacts loaded: {len(contacts)}")

    if not contacts:
        print("No contacts found.")
        return

    results = []
    context = None

    try:
        with sync_playwright() as p:
            print("\nLaunching Google Chrome...")

            context = p.chromium.launch_persistent_context(
                user_data_dir=str(PROFILE_DIR),
                channel="chrome",
                headless=False,
                viewport=None
            )

            if context.pages:
                page = context.pages[0]
            else:
                page = context.new_page()

            wait_for_whatsapp_login(page)

            print("\nStarting contact processing...")

            for index, contact in enumerate(
                contacts,
                start=1
            ):
                result = process_contact(
                    page,
                    contact,
                    index
                )

                results.append(result)

                save_json_report(results)
                save_excel_report(results)

                print("\nProgress saved.")

                if index < len(contacts):
                    human_pause(page, 2, 5)

            print("\nAll contacts processed.")
            page.wait_for_timeout(3000)

            context.close()
            context = None

    except KeyboardInterrupt:
        print("\nAutomation stopped manually.")

    except Exception as error:
        print("\nAUTOMATION ERROR")
        print(error)

    finally:
        if results:
            try:
                save_json_report(results)
                save_excel_report(results)
            except Exception as report_error:
                print(
                    "Could not save final report:",
                    report_error
                )

        if context is not None:
            try:
                context.close()
            except Exception:
                pass

    sent_count = sum(
        1 for result in results
        if result.get("status") == "SENT"
    )

    not_found_count = sum(
        1 for result in results
        if result.get("status") == "CONTACT_NOT_FOUND"
    )

    failed_count = sum(
        1 for result in results
        if result.get("status") == "FAILED"
    )

    print("\n" + "=" * 70)
    print("AUTOMATION COMPLETED")
    print("=" * 70)

    print(f"\nContacts processed: {len(results)}")
    print(f"Confirmed sent: {sent_count}")
    print(f"Contact not found: {not_found_count}")
    print(f"Failed: {failed_count}")

    if results:
        print(f"\nJSON report:\n{JSON_REPORT}")
        print(f"\nExcel report:\n{EXCEL_REPORT}")
        print(f"\nScreenshots folder:\n{SCREENSHOT_DIR}")


if __name__ == "__main__":
    main()
